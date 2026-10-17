"""Genera hoja/Asistencia_Docentes_UFPS.xlsx a partir de herramientas/datos_horario.py.
Uso: python3 herramientas/generar_hoja.py   (requiere: pip install openpyxl)"""
import sys, json, os
AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
from datos_horario import docentes
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from datetime import date

DIAS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado"]
AB = {"Lunes": "LUN", "Martes": "MAR", "Miércoles": "MIE", "Jueves": "JUE", "Viernes": "VIE", "Sábado": "SAB"}
CAT = {"TC": "Tiempo completo", "CT": "Cátedra", "OTC": "Ocasional TC"}
OUT = os.path.join(RAIZ, "hoja", "Asistencia_Docentes_UFPS.xlsx")

rows = []
for cd, nom, cat, mats in docentes:
    for cm, alu, nh, mat, ses in mats:
        for dia, ini, fin, sal in ses:
            rows.append(dict(id=f"{cd}-{cm}-{AB[dia]}-{ini.replace(':','')}", dia=dia, ini=ini, fin=fin, cm=cm, mat=mat, cd=cd, nom=nom, sal=sal, alu=alu))
rows.sort(key=lambda r: (DIAS.index(r["dia"]), r["ini"], r["sal"]))
json.dump(rows, open(os.path.join(RAIZ, "test", "fixtures", "horario.json"), "w"), ensure_ascii=False)

from openpyxl.drawing.image import Image as XLImage

# Identidad UFPS: rojo institucional + colores de la bandera
ROJO, NEGRO, NARANJA, VERDE = "E2031A", "111111", "FF9900", "66CC33"
MARCA = os.path.join(RAIZ, "marca")
F = "Arial"
HF = PatternFill("solid", fgColor=ROJO); HFONT = Font(name=F, bold=True, color="FFFFFF", size=10)
B = Font(name=F, size=10); BOLD = Font(name=F, bold=True, size=10); NOTE = Font(name=F, size=9, color="5B5E63")
INPUT = PatternFill("solid", fgColor="FFF4E0"); BLUE = Font(name=F, size=10, color="0000FF")
TITULO = lambda size=16: Font(name=F, bold=True, size=size, color=ROJO)

def logo(ws, archivo, celda, alto):
    img = XLImage(os.path.join(MARCA, archivo))
    img.width, img.height = round(img.width * alto / img.height), alto
    ws.add_image(img, celda)

def bandera(ws, fila, col_fin):
    """Tres franjas finas (negro, naranja, verde) como la bandera de la UFPS."""
    for k, color in enumerate((NEGRO, NARANJA, VERDE)):
        ws.row_dimensions[fila + k].height = 4
        for c in range(1, col_fin + 1):
            ws.cell(fila + k, c).fill = PatternFill("solid", fgColor=color)
th = Side(style="thin", color="BFBFBF"); BORD = Border(left=th, right=th, top=th, bottom=th)

def header(ws, cols, widths, r=1):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(r, i, c); x.font = HFONT; x.fill = HF; x.border = BORD
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[r].height = 30
    ws.freeze_panes = ws.cell(r + 1, 1)

def put(ws, r, vals, text_cols=()):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = B; c.border = BORD
        if j in text_cols: c.number_format = "@"

def lista(ws, rango, opciones):
    dv = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"'); ws.add_data_validation(dv); dv.add(rango)

wb = Workbook()

# ---------------- Instrucciones ----------------
wi = wb.active; wi.title = "Instrucciones"
wi.column_dimensions["A"].width = 12; wi.column_dimensions["B"].width = 108; wi.column_dimensions["C"].width = 14
wi.row_dimensions[1].height = 62
wi["B1"] = "Pulso – Asistencia docente · Departamento de Electricidad y Electrónica"; wi["B1"].font = TITULO(16)
wi["B1"].alignment = Alignment(vertical="center")
logo(wi, "ufps.png", "A1", 72)
logo(wi, "ing-electronica.png", "C1", 72)
bandera(wi, 2, 3)
wi.sheet_view.showGridLines = False
pasos = [
    ("CÓMO LO USA EL DOCENTE", None),
    ("1", "Recibe por correo su enlace personal y lo guarda en la pantalla de inicio del celular (una sola vez)."),
    ("2", "Al llegar al salón: se conecta al WiFi de la UFPS, abre el enlace y toca “Registrar asistencia”. La página ya muestra su clase, el salón y la hora."),
    ("POR QUÉ ES SEGURO SIN PANTALLAS NI PERSONAL", None),
    ("•", "Red WiFi: la página pasa por un Cloudflare Worker (gratis) que ve la IP real de la conexión. Un GPS falso cambia la ubicación, pero desde la casa la conexión no sale por la red de la UFPS. El celular no puede falsificar esa IP."),
    ("•", "GPS: 3 lecturas frescas, precisión mínima y radio del campus. Además detecta patrones de apps de GPS falso (coordenadas congeladas, precisión irreal, saltos imposibles, coordenadas idénticas a otro día)."),
    ("•", "Celular vinculado: el enlace queda atado al primer celular donde se abre. Si se reenvía el enlace a otra persona, no funciona."),
    ("•", "Horario: solo se puede registrar desde 15 min antes del inicio hasta el final de la clase que el docente tiene programada."),
    ("•", "Firma: el Apps Script solo acepta solicitudes firmadas por el Worker; no se puede saltar llamando directo al servidor."),
    ("•", "Foto en vivo: opcional (Config). Si la activas, queda la selfie en Drive como evidencia."),
    ("INSTALACIÓN DESDE GITHUB", None),
    ("1", "Sube este archivo a Google Drive → Abrir con Hojas de cálculo de Google → Archivo → Guardar como Hojas de cálculo de Google. Mantén la hoja PRIVADA."),
    ("2", "En la hoja: Extensiones → Apps Script → engranaje (Configuración del proyecto) → copia el “ID de la secuencia de comandos”. Ese es el SCRIPT_ID para GitHub."),
    ("3", "Sigue el README del repositorio: el workflow “Desplegar” de GitHub sube el código a Apps Script y a Cloudflare automáticamente."),
    ("4", "Después del primer despliegue: recarga esta hoja → menú Pulso · Asistencia → 1. Autorizar y verificar instalación (acepta los permisos)."),
    ("5", "En Config pega la “URL de la página (GitHub Pages)” que aparece en el resumen del workflow en GitHub, tu correo y las fechas. Llena la columna Correo en Docentes."),
    ("6", "Menú → 2. Generar enlaces → 3. Enviar enlaces → 4. Activar automatización."),
    ("•", "Cambios futuros al código: se hacen en GitHub. Cada cambio en la rama main se prueba y se despliega solo; la URL no cambia."),
    ("ACTIVAR LA RED WIFI (semanas 1–2)", None),
    ("•", "Empieza con Validación de red = ALERTA: registra y marca quién está fuera de la red, sin rechazar."),
    ("•", "Cada noche la hoja Redes se llena sola con las redes desde donde registran los docentes. Escribe SI en “Aprobada” solo en la red de la universidad (la que usan muchos docentes y no es Claro/Movistar/Tigo/WOM). Mejor aún: pide a Sistemas de la UFPS los rangos de IP públicos del WiFi y agrégalos como filas aprobadas."),
    ("•", "Cuando la red esté aprobada, cambia Validación de red a EXIGIR: desde ese momento solo se acepta el registro conectado al WiFi UFPS."),
    ("LÍMITES (para que lo tengas claro)", None),
    ("•", "Valida que el docente está en la universidad durante su horario, no el salón exacto: sin hardware (pantallas o balizas) ningún celular puede probar en qué salón está. Puedes acotar por bloque llenando lat/lng/radio en Salones."),
    ("•", "Si la UFPS tiene VPN para docentes, excluye su rango (no lo apruebes) para que no se pueda registrar desde la casa por VPN."),
    ("•", "Un docente podría entregarle su celular a otra persona en el campus. Para eso existe la foto opcional y la alerta de celular compartido."),
    ("•", "La ubicación, la IP y la foto son datos personales (Ley 1581 de 2012): informa a los docentes y obtén su autorización."),
    ("MANTENIMIENTO", None),
    ("•", "Docente cambió de celular o borró los datos del navegador: borra su celda en Docentes → Celular vinculado."),
    ("•", "Enlace filtrado: borra su celda Clave del enlace → menú 2 y 3 (el enlace anterior deja de funcionar)."),
    ("•", "Salón virtual o fuera del campus (TE000): Exige ubicación = NO en Salones (queda con alerta)."),
]
r = 6
for a, b in pasos:
    if b is None:
        wi.cell(r, 2, a).font = Font(name=F, bold=True, size=11, color=ROJO); r += 1; continue
    n = wi.cell(r, 1, a); n.font = Font(name=F, bold=True, size=10, color=ROJO); n.alignment = Alignment(horizontal="center", vertical="top")
    c = wi.cell(r, 2, b); c.font = B; c.alignment = Alignment(wrap_text=True, vertical="top")
    wi.row_dimensions[r].height = 15 * (1 + len(b) // 100)
    r += 1

# ---------------- Config ----------------
wc = wb.create_sheet("Config")
header(wc, ["Parámetro", "Valor", "Nota"], [42, 36, 84])
cfg = [
    ("URL de la página (GitHub Pages)", "", "Dirección de la página del docente (index.html). Aparece en el resumen del workflow Desplegar (ej. https://usuario.github.io/pulso-ufps/)."),
    ("URL del Worker", "", "Opcional: puente de Cloudflare (ej. https://pulso-ufps.tu-usuario.workers.dev). Se usa si no hay URL de GitHub Pages."),
    ("Correo del coordinador", "", "Alertas y reporte diario. Varios correos separados por coma."),
    ("Validación de red WiFi (NO/ALERTA/EXIGIR)", "ALERTA", "ALERTA las primeras semanas mientras apruebas la red en la hoja Redes; luego EXIGIR."),
    ("Latitud del campus", 7.898115, "UFPS sede principal, Av. Gran Colombia 12E-96 (Google Maps)."),
    ("Longitud del campus", -72.488703, ""),
    ("Radio permitido (m)", 300, "Distancia máxima al punto central (cada salón puede tener el suyo en Salones)."),
    ("Precisión GPS máxima (m)", 100, "Lecturas menos precisas se rechazan."),
    ("Lecturas GPS mínimas", 3, ""),
    ("Bloquear si cambia el celular (SI/NO)", "SI", "SI: el enlace solo funciona en el celular vinculado. NO: solo alerta."),
    ("Exigir foto en vivo (SI/NO)", "NO", "SI agrega una selfie como evidencia (guardada en Drive)."),
    ("Correo inmediato por alertas (SI/NO)", "SI", "Registro aceptado con algo sospechoso (GPS, red, celular)."),
    ("Minutos antes del inicio para registrar", 15, ""),
    ("Tolerancia de llegada (min)", 10, "Hasta cuántos minutos después del inicio cuenta como “A tiempo”."),
    ("Enviar alertas de no registro (SI/NO)", "SI", "Correo cuando una clase pasa la tolerancia sin registro."),
    ("Fecha inicio de control", date(2026, 10, 19), "Antes de esta fecha no se marcan inasistencias."),
    ("Fecha fin de control", "", "Último día de clases del semestre."),
]
for i, (k, v, n) in enumerate(cfg, 2):
    wc.cell(i, 1, k).font = BOLD
    c = wc.cell(i, 2, v); c.font = BLUE; c.fill = INPUT; c.border = BORD
    if k.startswith(("Latitud", "Longitud")): c.number_format = "0.000000"
    if k.startswith("Fecha"): c.number_format = "yyyy-mm-dd"
    wc.cell(i, 3, n).font = NOTE
    if "(SI/NO)" in k: lista(wc, f"B{i}", ["SI", "NO"])
    if "(NO/ALERTA/EXIGIR)" in k: lista(wc, f"B{i}", ["NO", "ALERTA", "EXIGIR"])

# ---------------- Horario ----------------
wh = wb.create_sheet("Horario")
header(wh, ["ID clase", "Día", "Hora inicio", "Hora fin", "Código materia", "Materia", "Código docente", "Docente", "Salón", "Alumnos"], [26, 11, 11, 10, 15, 46, 13, 42, 10, 9])
for i, x in enumerate(rows, 2):
    put(wh, i, [x["id"], x["dia"], x["ini"], x["fin"], x["cm"], x["mat"], x["cd"], x["nom"], x["sal"], x["alu"]], text_cols=(1, 3, 4, 5, 7))
wh.auto_filter.ref = f"A1:J{len(rows)+1}"

# ---------------- Docentes ----------------
wd = wb.create_sheet("Docentes")
header(wd, ["Código docente", "Docente", "Categoría", "Correo", "Activo (SI/NO)", "Clave del enlace (menú 2)", "Celular vinculado (automático)", "Enlace personal (menú 2) – PRIVADO"], [14, 42, 16, 32, 10, 36, 40, 80])
dl = sorted(docentes, key=lambda d: d[1])
for i, (cd, nom, cat, _) in enumerate(dl, 2):
    put(wd, i, [cd, nom, CAT[cat], "", "SI", "", "", ""], text_cols=(1, 6, 7))
    wd.cell(i, 4).fill = INPUT; wd.cell(i, 5).fill = INPUT
lista(wd, f"E2:E{len(dl)+1}", ["SI", "NO"])

# ---------------- Salones ----------------
ws = wb.create_sheet("Salones")
header(ws, ["Salón", "Exige ubicación y red (SI/NO)", "Latitud (opcional)", "Longitud (opcional)", "Radio (m, opcional)"], [10, 16, 13, 13, 12])
sals = sorted({x["sal"] for x in rows})
for i, sa in enumerate(sals, 2):
    put(ws, i, [sa, "NO" if sa == "TE000" else "SI", "", "", ""])
    for col in range(2, 6): ws.cell(i, col).fill = INPUT
lista(ws, f"B2:B{len(sals)+1}", ["SI", "NO"])
ws.cell(len(sals) + 3, 1, "Lat/Lng/Radio opcionales: para exigir un bloque concreto (ej. bloque LG). Vacío = se usa el campus de Config.").font = NOTE

# ---------------- Redes ----------------
wr = wb.create_sheet("Redes")
header(wr, ["Rango de red", "Organización / proveedor", "Docentes distintos", "Registros", "Última vez", "Aprobada (SI/NO/PENDIENTE)"], [24, 40, 12, 10, 12, 16])
lista(wr, "F2:F500", ["SI", "NO", "PENDIENTE"])
wr.cell(2, 8, "Se llena sola cada noche con las redes desde donde registran los docentes (GPS dentro del campus).").font = NOTE
wr.cell(3, 8, "Escribe SI solo en la red de la universidad. NO en operadores móviles (Claro, Movistar, Tigo, WOM) ni VPN.").font = NOTE
wr.cell(4, 8, "También puedes agregar a mano los rangos que te dé Sistemas (ej. 190.25.10.0/24) con Aprobada = SI.").font = NOTE
wr.column_dimensions["H"].width = 90

# ---------------- Festivos ----------------
wf = wb.create_sheet("Festivos")
header(wf, ["Fecha", "Descripción"], [14, 60])
for i, (d, t) in enumerate([(date(2026, 11, 2), "Día de Todos los Santos (trasladado)"), (date(2026, 11, 16), "Independencia de Cartagena (trasladado)"),
                            (date(2026, 12, 8), "Inmaculada Concepción"), (date(2026, 12, 25), "Navidad")], 2):
    put(wf, i, [d, t]); wf.cell(i, 1).number_format = "yyyy-mm-dd"
wf.cell(7, 1, "Agrega recesos o días sin clase. Verifica con el calendario académico oficial.").font = NOTE

# ---------------- Registros ----------------
wg = wb.create_sheet("Registros")
header(wg, ["Fecha", "Hora", "Código", "Docente", "Salón", "Latitud", "Longitud", "Precisión (m)", "Lecturas GPS", "Resultado", "Motivo", "ID clase",
            "Alertas", "Distancia (m)", "Celular", "IP", "Red", "Navegador"],
       [11, 10, 9, 36, 9, 11, 12, 10, 9, 11, 46, 26, 46, 10, 36, 18, 28, 36])

# ---------------- Asistencia ----------------
wa = wb.create_sheet("Asistencia")
header(wa, ["Fecha", "Día", "ID clase", "Hora inicio", "Hora fin", "Código materia", "Materia", "Código docente", "Docente", "Salón",
            "Hora registro", "Minutos de retraso", "Estado", "En red UFPS", "Distancia (m)", "Alertas", "Foto", "Revisión coordinador"],
       [11, 10, 26, 9, 9, 13, 40, 11, 38, 9, 11, 10, 12, 9, 10, 46, 26, 16])
wa.auto_filter.ref = "A1:R1"
lista(wa, "R2:R5000", ["Válida", "No válida"])
from openpyxl.formatting.rule import FormulaRule
for texto, fondo, letra in (("A tiempo", "E3F5D8", "2E7D17"), ("Tarde", "FFF4E0", "8A4B00"), ("No registró", "FDECEE", "B80215")):
    wa.conditional_formatting.add("M2:M5000", FormulaRule(formula=[f'$M2="{texto}"'], fill=PatternFill("solid", fgColor=fondo), font=Font(name=F, color=letra, bold=True)))

# ---------------- Resumen ----------------
wsr = wb.create_sheet("Resumen")
wsr["A1"] = "Cumplimiento acumulado por docente"; wsr["A1"].font = TITULO(14)
wsr["A2"] = "Se calcula desde Asistencia. Cumplimiento = (A tiempo + Tarde − revisadas No válidas) / clases controladas."; wsr["A2"].font = NOTE
wsr.row_dimensions[1].height = 40
header(wsr, ["Código docente", "Docente", "Clases controladas", "A tiempo", "Tarde", "No registró", "Fuera de red UFPS", "Cumplimiento",
             "Retraso promedio (min)", "Con alerta", "Revisadas no válidas"], [14, 42, 12, 10, 10, 11, 12, 13, 13, 10, 12], r=4)
for i, (cd, nom, cat, _) in enumerate(dl, 5):
    n = f"$B{i}"
    put(wsr, i, [cd, nom,
        f'=COUNTIFS(Asistencia!$I:$I,{n})',
        f'=COUNTIFS(Asistencia!$I:$I,{n},Asistencia!$M:$M,"A tiempo")',
        f'=COUNTIFS(Asistencia!$I:$I,{n},Asistencia!$M:$M,"Tarde")',
        f'=COUNTIFS(Asistencia!$I:$I,{n},Asistencia!$M:$M,"No registró")',
        f'=COUNTIFS(Asistencia!$I:$I,{n},Asistencia!$N:$N,"NO")',
        f'=IF(C{i}=0,"",(D{i}+E{i}-K{i})/C{i})',
        f'=IFERROR(AVERAGEIFS(Asistencia!$L:$L,Asistencia!$I:$I,{n},Asistencia!$M:$M,"Tarde"),"")',
        f'=COUNTIFS(Asistencia!$I:$I,{n},Asistencia!$P:$P,"?*")',
        f'=COUNTIFS(Asistencia!$I:$I,{n},Asistencia!$R:$R,"No válida")'], text_cols=(1,))
    wsr.cell(i, 8).number_format = "0%"; wsr.cell(i, 9).number_format = "0.0"
tot = len(dl) + 5
wsr.cell(tot, 2, "TOTAL").font = BOLD
for col in "CDEFGJK":
    c = wsr[f"{col}{tot}"]; c.value = f"=SUM({col}5:{col}{tot-1})"; c.font = BOLD; c.border = BORD
c = wsr[f"H{tot}"]; c.value = f'=IF(C{tot}=0,"",(D{tot}+E{tot}-K{tot})/C{tot})'; c.font = BOLD; c.number_format = "0%"; c.border = BORD
wsr.freeze_panes = "C5"
logo(wsr, "ufps.png", "L1", 48)
logo(wsr, "acreditado.png", "M1", 48)
wsr.column_dimensions["L"].width = 9
# Colores de semáforo con la paleta UFPS
from openpyxl.formatting.rule import CellIsRule
rng = f"H5:H{tot}"
wsr.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["0.9"], fill=PatternFill("solid", fgColor="E3F5D8"), font=Font(name=F, color="2E7D17", bold=True)))
wsr.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["0.75", "0.8999"], fill=PatternFill("solid", fgColor="FFF4E0"), font=Font(name=F, color="8A4B00", bold=True)))
wsr.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0.75"], fill=PatternFill("solid", fgColor="FDECEE"), font=Font(name=F, color="B80215", bold=True)))

for nombre, color in (("Instrucciones", ROJO), ("Config", NARANJA), ("Horario", NEGRO), ("Docentes", NEGRO), ("Salones", NEGRO),
                      ("Redes", NARANJA), ("Festivos", NEGRO), ("Registros", "909398"), ("Asistencia", VERDE), ("Resumen", VERDE)):
    wb[nombre].sheet_properties.tabColor = color
wb.save(OUT)
print(len(rows), "clases,", len(dl), "docentes,", len(sals), "salones")
